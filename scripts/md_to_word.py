from pathlib import Path
import sys
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Tử vi Mr Hoàng" / "THANG-08-AM-BINH-NGO-2026.md"
OUTPUT = ROOT / "tuvi-thang-08-binh-ngo-2026.docx"

NAVY = "17365D"
BLUE = "1F4E79"
RED = "9C0006"
GREEN = "216E39"
GOLD = "9C6500"
GRAY = "666666"
LIGHT_BLUE = "D9EAF7"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
LIGHT_GOLD = "FFF2CC"


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def shade_paragraph(paragraph, fill):
    properties = paragraph._p.get_or_add_pPr()
    element = properties.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        properties.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_inline(paragraph, text, color=None):
    tokens = re.split(r"(\*\*.*?\*\*)", text)
    for token in tokens:
        if not token:
            continue
        bold = token.startswith("**") and token.endswith("**")
        value = token[2:-2] if bold else token
        run = paragraph.add_run(value)
        run.bold = bold
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, color=None, style=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.12
    add_inline(paragraph, text, color)
    return paragraph


def add_bullets(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    add_inline(paragraph, text)


def add_table(doc, lines):
    rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    if len(rows) < 2:
        return
    header = rows[0]
    data = [row for row in rows[2:] if len(row) == len(header)]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, value in enumerate(header):
        set_cell_text(table.rows[0].cells[index], value, bold=True, color="FFFFFF", size=8)
        shade(table.rows[0].cells[index], NAVY)
    for row in data:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            color = None
            bold = index == 0
            if "Rất tốt" in value:
                shade(cells[index], LIGHT_GREEN)
                color = GREEN
            elif "Cẩn trọng" in value:
                shade(cells[index], LIGHT_RED)
                color = RED
            set_cell_text(cells[index], value, bold=bold, color=color, size=8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def make_doc():
    source_path = Path(sys.argv[1]) if len(sys.argv) > 1 else SOURCE
    output_path = Path(sys.argv[2]) if len(sys.argv) > 2 else OUTPUT
    source = source_path.read_text(encoding="utf-8")
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    for name, size, color in [("Heading 1", 18, NAVY), ("Heading 2", 14, BLUE), ("Heading 3", 11, BLUE)]:
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")

    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line[2:])
            run.bold = True
            run.font.name = "Aptos Display"
            run.font.size = Pt(22)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        elif line.startswith("## "):
            text = line[3:]
            color = RED if any(word in text for word in ["Rủi ro", "Kết luận"]) else NAVY
            paragraph = document.add_paragraph(style="Heading 1")
            run = paragraph.add_run(text)
            run.font.color.rgb = RGBColor.from_string(color)
        elif line.startswith("### "):
            text = line[4:]
            color = RED if any(word in text for word in ["Rủi ro", "Tránh", "Cảnh báo"]) else (GREEN if any(word in text for word in ["Cơ hội", "Nên", "Tài Bạch"]) else BLUE)
            paragraph = document.add_paragraph(style="Heading 2")
            run = paragraph.add_run(text)
            run.font.color.rgb = RGBColor.from_string(color)
        elif line.startswith("#### "):
            document.add_paragraph(line[5:], style="Heading 3")
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_after = Pt(8)
            shade_paragraph(paragraph, LIGHT_GOLD)
            add_inline(paragraph, line[2:])
        elif line.startswith("- "):
            add_bullets(document, line[2:])
        elif line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        elif line == "---":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(3)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "8")
            bottom.set(qn("w:color"), "D9E2F3")
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
        else:
            add_body(document, line)
        index += 1

    document.core_properties.title = "Tử vi tháng 08 âm lịch Bính Ngọ 2026 - Hoàng"
    document.core_properties.subject = "Luận giải chi tiết theo ngày và 12 cung"
    document.core_properties.author = "Tử Vi AI"
    document.save(output_path)
    print(output_path)


if __name__ == "__main__":
    make_doc()